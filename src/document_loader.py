"""
Document Loader Module
Loads documents across various formats (.docx, .png) into a unified representation.
"""

import os
from dataclasses import dataclass
import docx
from src.ocr_processor import extract_text_from_image

@dataclass
class Document:
    """Unified container for ingested document data."""
    file_name: str
    file_type: str
    file_path: str
    extracted_text: str

class DocumentLoader:
    """Unified document loader for DOCX and image-based scanned documents."""

    @staticmethod
    def extract_from_docx(file_path: str) -> str:
        """Extracts text content from a DOCX file including paragraphs and tables."""
        doc = docx.Document(file_path)
        text_blocks = []
        
        # Extract paragraph text
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_blocks.append(paragraph.text.strip())
                
        # Extract table text if any
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_blocks.append(" | ".join(row_text))
                    
        return "\n".join(text_blocks)

    @classmethod
    def load_document(cls, file_path: str) -> Document:
        """
        Loads a single document and extracts its textual content.
        
        Args:
            file_path (str): Path to the target document (.docx or .png).
            
        Returns:
            Document: Structured document object containing metadata and raw text.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        base_name = os.path.basename(file_path)
        # Handle composite extensions such as .pdf.docx
        if base_name.endswith(".pdf.docx"):
            doc_id = base_name[:-9]
            ext = ".docx"
        else:
            doc_id, ext = os.path.splitext(base_name)
            ext = ext.lower()

        if ext == ".docx":
            text = cls.extract_from_docx(file_path)
            file_type = "DOCX"
        elif ext in [".png", ".jpg", ".jpeg"]:
            text = extract_text_from_image(file_path)
            file_type = "PNG/Scanned Image"
        else:
            raise ValueError(f"Unsupported file format: {ext} for file {file_path}")

        return Document(
            file_name=doc_id,
            file_type=file_type,
            file_path=file_path,
            extracted_text=text
        )
